"""
Aplikasi Web SIMONKIN Internal
----------------------------------------
Fitur:
  - Login & hak akses (Admin = akses penuh, User = hanya mengerjakan tugas/assignment)
  - Data pegawai
  - Jadwal rapat + disposisi pegawai
  - Laporan status kinerja per pegawai (EFEKTIF wajib file bukti, TIDAK EFEKTIF file otomatis kosong)
  - Kalender, event & reminder (notifikasi di browser)
  - Rekap kinerja + ekspor Excel
  - Pengaturan sistem sederhana

Jalankan:  python app.py   (default http://0.0.0.0:5000)
Login awal: admin / admin123  (SEGERA ganti password setelah login)
"""
import calendar
import io
import os
import secrets
import sqlite3
import time
import uuid
from datetime import date, datetime, timedelta
from functools import wraps

from flask import (Flask, abort, flash, g, jsonify, redirect, render_template,
                   request, send_file, send_from_directory, session, url_for)
from markupsafe import Markup
from werkzeug.security import check_password_hash, generate_password_hash

# ---------------------------------------------------------------------------
# Konfigurasi
# ---------------------------------------------------------------------------
os.environ.setdefault("TZ", os.environ.get("APP_TZ", "Asia/Jakarta"))
if hasattr(time, "tzset"):
    time.tzset()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR", os.path.join(BASE_DIR, "data"))
DB_PATH = os.path.join(DATA_DIR, "monitoring.db")
UPLOAD_DIR = os.path.join(DATA_DIR, "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

ALLOWED_EXT = {"pdf", "png", "jpg", "jpeg", "doc", "docx", "xls", "xlsx", "ppt", "pptx"}
MAX_UPLOAD_MB = int(os.environ.get("MAX_UPLOAD_MB", "10"))
FMT = "%Y-%m-%dT%H:%M"
STATUS_LIST = ("EFEKTIF", "TIDAK EFEKTIF")


def _secret_key():
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    path = os.path.join(DATA_DIR, "secret.key")
    if not os.path.exists(path):
        with open(path, "w") as fh:
            fh.write(secrets.token_hex(32))
    with open(path) as fh:
        return fh.read().strip()


app = Flask(__name__)
app.config.update(
    SECRET_KEY=_secret_key(),
    MAX_CONTENT_LENGTH=MAX_UPLOAD_MB * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
)
if os.environ.get("COOKIE_CROSS_SITE") == "1":
    # Hanya untuk pratinjau dalam iframe (HTTPS). Tidak diperlukan untuk pemakaian normal di kantor.
    app.config.update(SESSION_COOKIE_SAMESITE="None", SESSION_COOKIE_SECURE=True, SESSION_COOKIE_PARTITIONED=True)

HARI = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
BULAN = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
         "Agustus", "September", "Oktober", "November", "Desember"]

DEFAULT_SETTINGS = {
    "nama_aplikasi": "SIMONKIN",
    "nama_instansi": "Bagian Umum",
    "pengingat_rapat_menit": "60",
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS pegawai (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    nip TEXT, nama TEXT NOT NULL, jabatan TEXT, unit TEXT,
    no_hp TEXT, email TEXT,
    aktif INTEGER NOT NULL DEFAULT 1,
    created_at TEXT
);
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT UNIQUE NOT NULL,
    password_hash TEXT NOT NULL,
    role TEXT NOT NULL CHECK (role IN ('admin','user')),
    pegawai_id INTEGER REFERENCES pegawai(id) ON DELETE SET NULL,
    aktif INTEGER NOT NULL DEFAULT 1,
    last_login TEXT, created_at TEXT
);
CREATE TABLE IF NOT EXISTS rapat (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    no_dispo_kabag TEXT, no_surat_undangan TEXT, asal_undangan TEXT,
    perihal TEXT NOT NULL, pimpinan_rapat TEXT, tempat TEXT,
    waktu TEXT NOT NULL, keterangan TEXT,
    created_by INTEGER, created_at TEXT, updated_at TEXT
);
CREATE TABLE IF NOT EXISTS disposisi (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    rapat_id INTEGER NOT NULL REFERENCES rapat(id) ON DELETE CASCADE,
    pegawai_id INTEGER NOT NULL REFERENCES pegawai(id),
    status TEXT CHECK (status IN ('EFEKTIF','TIDAK EFEKTIF')),
    file_path TEXT, file_name TEXT, catatan TEXT,
    dilaporkan_at TEXT, dilaporkan_oleh INTEGER,
    UNIQUE (rapat_id, pegawai_id)
);
CREATE TABLE IF NOT EXISTS events (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    judul TEXT NOT NULL, deskripsi TEXT,
    jenis TEXT NOT NULL DEFAULT 'event' CHECK (jenis IN ('event','reminder')),
    mulai TEXT NOT NULL, selesai TEXT,
    ingatkan_menit INTEGER NOT NULL DEFAULT 0,
    untuk_pegawai_id INTEGER REFERENCES pegawai(id) ON DELETE CASCADE,
    created_by INTEGER, created_at TEXT
);
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);
CREATE INDEX IF NOT EXISTS idx_rapat_waktu ON rapat(waktu);
CREATE INDEX IF NOT EXISTS idx_dispo_pegawai ON disposisi(pegawai_id);
CREATE INDEX IF NOT EXISTS idx_events_mulai ON events(mulai);
"""

# ---------------------------------------------------------------------------
# Database helpers
# ---------------------------------------------------------------------------

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def query(sql, args=(), one=False):
    cur = get_db().execute(sql, args)
    rows = cur.fetchall()
    cur.close()
    return (rows[0] if rows else None) if one else rows


def execute(sql, args=()):
    db = get_db()
    cur = db.execute(sql, args)
    db.commit()
    return cur.lastrowid


def now_str():
    return datetime.now().strftime(FMT)


def parse_dt(value):
    if not value:
        return None
    for fmt in (FMT, "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(value[:16] if "T" in value or " " in value else value[:10], fmt)
        except (ValueError, TypeError):
            continue
    return None


def init_db():
    os.makedirs(DATA_DIR, exist_ok=True)
    fresh = not os.path.exists(DB_PATH)
    db = sqlite3.connect(DB_PATH)
    db.executescript(SCHEMA)
    for k, v in DEFAULT_SETTINGS.items():
        db.execute("INSERT OR IGNORE INTO settings(key, value) VALUES (?, ?)", (k, v))
    if not db.execute("SELECT 1 FROM users WHERE role='admin' LIMIT 1").fetchone():
        db.execute(
            "INSERT INTO users(username, password_hash, role, created_at) VALUES (?,?,?,?)",
            ("admin", generate_password_hash("admin123"), "admin", now_str()),
        )
    db.commit()
    if fresh and os.environ.get("SEED_DEMO", "1") == "1":
        seed_demo(db)
    db.close()


def seed_demo(db):
    """Data contoh agar aplikasi langsung bisa dicoba. Matikan dengan SEED_DEMO=0."""
    ts = now_str()
    pegawai = [
        ("198501012010011001", "Budi Santoso", "Analis Kebijakan", "Subbag Umum", "081200000001", "budi@contoh.go.id"),
        ("199002022015032002", "Siti Rahmawati", "Pranata Humas", "Subbag Humas", "081200000002", "siti@contoh.go.id"),
        ("198703032012121003", "Ahmad Fauzi", "Pengelola Keuangan", "Subbag Keuangan", "081200000003", "ahmad@contoh.go.id"),
        ("199204042019022004", "Dewi Lestari", "Arsiparis", "Subbag Umum", "081200000004", "dewi@contoh.go.id"),
        ("198805052014031005", "Rizky Pratama", "Pranata Komputer", "Subbag Program", "081200000005", "rizky@contoh.go.id"),
    ]
    ids = []
    for p in pegawai:
        cur = db.execute(
            "INSERT INTO pegawai(nip,nama,jabatan,unit,no_hp,email,created_at) VALUES (?,?,?,?,?,?,?)", (*p, ts))
        ids.append(cur.lastrowid)
    db.execute("INSERT INTO users(username,password_hash,role,pegawai_id,created_at) VALUES (?,?,?,?,?)",
               ("budi", generate_password_hash("user123"), "user", ids[0], ts))
    db.execute("INSERT INTO users(username,password_hash,role,pegawai_id,created_at) VALUES (?,?,?,?,?)",
               ("siti", generate_password_hash("user123"), "user", ids[1], ts))

    now = datetime.now().replace(second=0, microsecond=0)
    rapat = [
        ("045/DISPO/KBG/X/2026", "005/123/Bappeda", "Bappeda Kabupaten",
         "Rapat Koordinasi Penyusunan RKPD Tahun 2027", "Sekretaris Daerah", "Ruang Rapat Bappeda Lt. 2",
         (now - timedelta(days=3)).replace(hour=9, minute=0), "Membawa bahan paparan", [ids[0], ids[2]],
         {ids[0]: ("TIDAK EFEKTIF", "Rapat ditunda oleh penyelenggara.")}),
        ("046/DISPO/KBG/X/2026", "005/456/Diskominfo", "Dinas Kominfo",
         "Sosialisasi Keamanan Informasi dan Tata Kelola SPBE", "Kepala Dinas Kominfo", "Aula Diskominfo",
         (now - timedelta(days=1)).replace(hour=13, minute=30), "", [ids[1], ids[4]], {}),
        ("047/DISPO/KBG/X/2026", "005/789/BKPSDM", "BKPSDM",
         "Rapat Evaluasi Kinerja Pegawai Triwulan III", "Kepala BKPSDM", "Ruang Rapat Utama",
         (now + timedelta(days=2)).replace(hour=10, minute=0), "Pakaian dinas harian", [ids[0], ids[1], ids[3]], {}),
        ("048/DISPO/KBG/X/2026", "005/321/Setda", "Sekretariat Daerah",
         "Persiapan Peringatan Hari Jadi Kabupaten", "Asisten Administrasi Umum", "Ruang Rapat Setda",
         (now + timedelta(days=5)).replace(hour=8, minute=30), "", [ids[0], ids[4]], {}),
    ]
    for r in rapat:
        cur = db.execute(
            """INSERT INTO rapat(no_dispo_kabag,no_surat_undangan,asal_undangan,perihal,pimpinan_rapat,tempat,
               waktu,keterangan,created_by,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,1,?,?)""",
            (*r[:6], r[6].strftime(FMT), r[7], ts, ts))
        rid = cur.lastrowid
        for pid in r[8]:
            st = r[9].get(pid)
            db.execute(
                "INSERT INTO disposisi(rapat_id,pegawai_id,status,catatan,dilaporkan_at,dilaporkan_oleh) VALUES (?,?,?,?,?,?)",
                (rid, pid, st[0] if st else None, st[1] if st else None, ts if st else None, 1 if st else None))
    db.execute("INSERT INTO events(judul,deskripsi,jenis,mulai,selesai,ingatkan_menit,created_by,created_at) VALUES (?,?,?,?,?,?,1,?)",
               ("Batas pengumpulan laporan bulanan", "Seluruh pegawai mengumpulkan laporan kinerja bulanan.",
                "reminder", (now + timedelta(days=1)).replace(hour=15, minute=0).strftime(FMT), "", 1440, ts))
    db.execute("INSERT INTO events(judul,deskripsi,jenis,mulai,selesai,ingatkan_menit,created_by,created_at) VALUES (?,?,?,?,?,?,1,?)",
               ("Apel Pagi Gabungan", "Halaman kantor", "event",
                (now + timedelta(days=4)).replace(hour=7, minute=30).strftime(FMT),
                (now + timedelta(days=4)).replace(hour=8, minute=30).strftime(FMT), 60, ts))
    db.commit()


# ---------------------------------------------------------------------------
# Auth, CSRF, context
# ---------------------------------------------------------------------------

def csrf_token():
    if "_csrf" not in session:
        session["_csrf"] = secrets.token_hex(16)
    return session["_csrf"]


@app.before_request
def before():
    g.settings = {r["key"]: r["value"] for r in query("SELECT key, value FROM settings")}
    for k, v in DEFAULT_SETTINGS.items():
        g.settings.setdefault(k, v)
    g.user = None
    uid = session.get("uid")
    if uid:
        g.user = query(
            """SELECT u.*, p.nama AS pegawai_nama, p.jabatan AS pegawai_jabatan
               FROM users u LEFT JOIN pegawai p ON p.id = u.pegawai_id
               WHERE u.id = ? AND u.aktif = 1""", (uid,), one=True)
        if g.user is None:
            session.clear()
    if request.method == "POST":
        token = request.form.get("_csrf") or request.headers.get("X-CSRF-Token")
        if not token or token != session.get("_csrf"):
            abort(400, "Token keamanan tidak valid. Muat ulang halaman lalu coba lagi.")


def is_admin():
    return bool(g.user and g.user["role"] == "admin")


def login_required(f):
    @wraps(f)
    def wrapper(*a, **kw):
        if not g.user:
            return redirect(url_for("login", next=request.full_path))
        return f(*a, **kw)
    return wrapper


def admin_required(f):
    @wraps(f)
    def wrapper(*a, **kw):
        if not g.user:
            return redirect(url_for("login", next=request.full_path))
        if g.user["role"] != "admin":
            abort(403)
        return f(*a, **kw)
    return wrapper


ICONS = {
    "home": '<path d="M3 11l9-8 9 8"/><path d="M5 10v10h14V10"/>',
    "rapat": '<rect x="3" y="4" width="18" height="16" rx="2"/><path d="M3 9h18M8 2v4M16 2v4"/>',
    "tugas": '<path d="M9 11l3 3 8-8"/><path d="M20 12v7a2 2 0 0 1-2 2H6a2 2 0 0 1-2-2V5a2 2 0 0 1 2-2h9"/>',
    "kalender": '<rect x="3" y="4" width="18" height="17" rx="2"/><path d="M3 10h18M8 2v4M16 2v4M8 14h2M14 14h2M8 17h2"/>',
    "bell": '<path d="M18 8a6 6 0 0 0-12 0c0 7-3 9-3 9h18s-3-2-3-9"/><path d="M13.7 21a2 2 0 0 1-3.4 0"/>',
    "pegawai": '<circle cx="9" cy="7" r="4"/><path d="M2 21v-2a4 4 0 0 1 4-4h6a4 4 0 0 1 4 4v2"/><path d="M16 3.1a4 4 0 0 1 0 7.8M22 21v-2a4 4 0 0 0-3-3.9"/>',
    "kunci": '<rect x="3" y="11" width="18" height="11" rx="2"/><path d="M7 11V7a5 5 0 0 1 10 0v4"/>',
    "chart": '<path d="M3 3v18h18"/><path d="M7 15l4-4 3 3 6-6"/>',
    "gear": '<circle cx="12" cy="12" r="3"/><path d="M19.4 15a1.7 1.7 0 0 0 .3 1.8l.1.1a2 2 0 1 1-2.8 2.8l-.1-.1a1.7 1.7 0 0 0-1.8-.3 1.7 1.7 0 0 0-1 1.5V21a2 2 0 1 1-4 0v-.1a1.7 1.7 0 0 0-1.1-1.5 1.7 1.7 0 0 0-1.8.3l-.1.1a2 2 0 1 1-2.8-2.8l.1-.1a1.7 1.7 0 0 0 .3-1.8 1.7 1.7 0 0 0-1.5-1H3a2 2 0 1 1 0-4h.1a1.7 1.7 0 0 0 1.5-1.1 1.7 1.7 0 0 0-.3-1.8l-.1-.1a2 2 0 1 1 2.8-2.8l.1.1a1.7 1.7 0 0 0 1.8.3H9a1.7 1.7 0 0 0 1-1.5V3a2 2 0 1 1 4 0v.1a1.7 1.7 0 0 0 1 1.5 1.7 1.7 0 0 0 1.8-.3l.1-.1a2 2 0 1 1 2.8 2.8l-.1.1a1.7 1.7 0 0 0-.3 1.8V9a1.7 1.7 0 0 0 1.5 1H21a2 2 0 1 1 0 4h-.1a1.7 1.7 0 0 0-1.5 1z"/>',
    "logout": '<path d="M9 21H5a2 2 0 0 1-2-2V5a2 2 0 0 1 2-2h4M16 17l5-5-5-5M21 12H9"/>',
    "plus": '<path d="M12 5v14M5 12h14"/>',
    "edit": '<path d="M12 20h9"/><path d="M16.5 3.5a2.1 2.1 0 0 1 3 3L7 19l-4 1 1-4z"/>',
    "trash": '<path d="M3 6h18M8 6V4h8v2M19 6l-1 14H6L5 6"/>',
    "eye": '<path d="M1 12s4-8 11-8 11 8 11 8-4 8-11 8S1 12 1 12z"/><circle cx="12" cy="12" r="3"/>',
    "file": '<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><path d="M14 2v6h6"/>',
    "download": '<path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4M7 10l5 5 5-5M12 15V3"/>',
    "menu": '<path d="M3 6h18M3 12h18M3 18h18"/>',
    "clock": '<circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 2"/>',
    "pin": '<path d="M21 10c0 7-9 13-9 13S3 17 3 10a9 9 0 0 1 18 0z"/><circle cx="12" cy="10" r="3"/>',
    "send": '<path d="M22 2L11 13M22 2l-7 20-4-9-9-4z"/>',
    "chev-l": '<path d="M15 18l-6-6 6-6"/>',
    "chev-r": '<path d="M9 18l6-6-6-6"/>',
    "user": '<circle cx="12" cy="8" r="4"/><path d="M4 21v-1a6 6 0 0 1 6-6h4a6 6 0 0 1 6 6v1"/>',
}


def icon(name, size=18):
    return Markup(
        f'<svg class="ic" width="{size}" height="{size}" viewBox="0 0 24 24" fill="none" stroke="currentColor" '
        f'stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">{ICONS.get(name, "")}</svg>')


app.jinja_env.globals.update(csrf_token=csrf_token, icon=icon)


@app.context_processor
def inject():
    belum = 0
    if g.get("user") and g.user["pegawai_id"]:
        row = query("SELECT COUNT(*) c FROM disposisi WHERE pegawai_id=? AND status IS NULL",
                    (g.user["pegawai_id"],), one=True)
        belum = row["c"] if row else 0
    return dict(settings=g.get("settings", DEFAULT_SETTINGS), current_user=g.get("user"),
                csrf_token=csrf_token, icon=icon, is_admin=is_admin(), badge_tugas=belum,
                now=now_str(), allowed_ext=", ".join(sorted(ALLOWED_EXT)), max_mb=MAX_UPLOAD_MB)


# ---------------------------------------------------------------------------
# Template filters (format tanggal Indonesia)
# ---------------------------------------------------------------------------

@app.template_filter("tgl")
def f_tgl(value, hari=True):
    dt = parse_dt(value)
    if not dt:
        return "-"
    s = f"{dt.day} {BULAN[dt.month]} {dt.year}"
    return f"{HARI[dt.weekday()]}, {s}" if hari else s


@app.template_filter("jam")
def f_jam(value):
    dt = parse_dt(value)
    return dt.strftime("%H.%M") if dt and len(value) > 10 else "-"


@app.template_filter("tgljam")
def f_tgljam(value):
    dt = parse_dt(value)
    if not dt:
        return "-"
    return f"{dt.day} {BULAN[dt.month][:3]} {dt.year}, {dt.strftime('%H.%M')}"


@app.template_filter("relatif")
def f_relatif(value):
    dt = parse_dt(value)
    if not dt:
        return ""
    diff = dt - datetime.now()
    mins = int(diff.total_seconds() // 60)
    if -60 < mins < 0:
        return "sedang berlangsung"
    if mins < 0:
        d = abs(mins) // 1440
        return "hari ini" if d == 0 else f"{d} hari lalu"
    if mins < 60:
        return f"{mins} menit lagi"
    if mins < 1440:
        return f"{mins // 60} jam lagi"
    return f"{mins // 1440} hari lagi"


@app.template_filter("bulan_nama")
def f_bulan(ym):
    try:
        y, m = ym.split("-")
        return f"{BULAN[int(m)]} {y}"
    except Exception:
        return ym


# ---------------------------------------------------------------------------
# Error handlers
# ---------------------------------------------------------------------------

@app.errorhandler(403)
def e403(_e):
    return render_template("error.html", code=403, pesan="Anda tidak memiliki akses ke halaman ini."), 403


@app.errorhandler(404)
def e404(_e):
    return render_template("error.html", code=404, pesan="Halaman atau data tidak ditemukan."), 404


@app.errorhandler(400)
def e400(e):
    return render_template("error.html", code=400, pesan=getattr(e, "description", "Permintaan tidak valid.")), 400


@app.errorhandler(413)
def e413(_e):
    flash(f"Ukuran file terlalu besar. Maksimal {MAX_UPLOAD_MB} MB.", "error")
    return redirect(request.referrer or url_for("dashboard"))


# ---------------------------------------------------------------------------
# Helpers data
# ---------------------------------------------------------------------------

def attach_dispo(rows):
    rows = [dict(r) for r in rows]
    if not rows:
        return rows
    ids = [r["id"] for r in rows]
    mapping = {}
    for i in range(0, len(ids), 500):
        chunk = ids[i:i + 500]
        qs = ",".join("?" * len(chunk))
        for d in query(f"""SELECT d.*, p.nama, p.jabatan, p.nip FROM disposisi d
                           JOIN pegawai p ON p.id = d.pegawai_id
                           WHERE d.rapat_id IN ({qs}) ORDER BY p.nama""", chunk):
            mapping.setdefault(d["rapat_id"], []).append(dict(d))
    for r in rows:
        r["dispo"] = mapping.get(r["id"], [])
        r["jml_belum"] = sum(1 for d in r["dispo"] if not d["status"])
    return rows


def visible_events_sql():
    """Filter event: admin melihat semua; user melihat event umum + event untuk dirinya."""
    if is_admin():
        return "1=1", []
    return "(e.untuk_pegawai_id IS NULL OR e.untuk_pegawai_id = ?)", [g.user["pegawai_id"] or -1]


def remove_file(path):
    if path:
        full = os.path.join(UPLOAD_DIR, os.path.basename(path))
        if os.path.exists(full):
            try:
                os.remove(full)
            except OSError:
                pass


def can_access_dispo(d):
    return is_admin() or (g.user["pegawai_id"] and d["pegawai_id"] == g.user["pegawai_id"])


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------

@app.route("/login", methods=["GET", "POST"])
def login():
    if g.user:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        u = query("SELECT * FROM users WHERE lower(username)=?", (username,), one=True)
        if u and u["aktif"] and check_password_hash(u["password_hash"], password):
            token = session.get("_csrf")
            session.clear()
            session["_csrf"] = token or secrets.token_hex(16)
            session["uid"] = u["id"]
            session.permanent = True
            execute("UPDATE users SET last_login=? WHERE id=?", (now_str(), u["id"]))
            nxt = request.args.get("next", "")
            if not nxt.startswith("/") or nxt.startswith("//"):
                nxt = url_for("dashboard")
            return redirect(nxt)
        if u and not u["aktif"]:
            flash("Akun Anda dinonaktifkan. Hubungi admin.", "error")
        else:
            flash("Username atau password salah.", "error")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    flash("Anda telah keluar.", "info")
    return redirect(url_for("login"))


@app.route("/akun/password", methods=["GET", "POST"])
@login_required
def ganti_password():
    if request.method == "POST":
        lama = request.form.get("lama", "")
        baru = request.form.get("baru", "")
        ulang = request.form.get("ulang", "")
        if not check_password_hash(g.user["password_hash"], lama):
            flash("Password lama salah.", "error")
        elif len(baru) < 6:
            flash("Password baru minimal 6 karakter.", "error")
        elif baru != ulang:
            flash("Konfirmasi password tidak sama.", "error")
        else:
            execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(baru), g.user["id"]))
            flash("Password berhasil diubah.", "success")
            return redirect(url_for("dashboard"))
    return render_template("password.html")


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

@app.route("/")
@login_required
def dashboard():
    now = now_str()
    week = (datetime.now() + timedelta(days=7)).strftime(FMT)
    today = date.today().isoformat()
    bulan = date.today().strftime("%Y-%m")
    ev_where, ev_args = visible_events_sql()
    events = query(f"""SELECT e.*, p.nama AS untuk_nama FROM events e LEFT JOIN pegawai p ON p.id=e.untuk_pegawai_id
                       WHERE {ev_where} AND substr(coalesce(nullif(e.selesai,''), e.mulai),1,10) >= ?
                       ORDER BY e.mulai LIMIT 6""", (*ev_args, today))
    ctx = dict(events=events, bulan=bulan)
    if is_admin():
        ctx["stats"] = query("""
            SELECT COUNT(DISTINCT r.id) rapat, COUNT(d.id) dispo,
                   COALESCE(SUM(d.status='EFEKTIF'),0) ef, COALESCE(SUM(d.status='TIDAK EFEKTIF'),0) tef,
                   COALESCE(SUM(d.id IS NOT NULL AND d.status IS NULL),0) belum
            FROM rapat r LEFT JOIN disposisi d ON d.rapat_id = r.id
            WHERE substr(r.waktu,1,7) = ?""", (bulan,), one=True)
        ctx["upcoming"] = attach_dispo(query(
            "SELECT * FROM rapat WHERE waktu >= ? AND waktu <= ? ORDER BY waktu LIMIT 8", (now[:10], week)))
        ctx["perlu"] = query("""
            SELECT d.id, d.rapat_id, p.nama, r.perihal, r.waktu FROM disposisi d
            JOIN rapat r ON r.id=d.rapat_id JOIN pegawai p ON p.id=d.pegawai_id
            WHERE d.status IS NULL AND r.waktu < ? ORDER BY r.waktu DESC LIMIT 8""", (now,))
        ctx["total_pegawai"] = query("SELECT COUNT(*) c FROM pegawai WHERE aktif=1", one=True)["c"]
    else:
        pid = g.user["pegawai_id"] or -1
        ctx["stats"] = query("""
            SELECT COUNT(*) total, COALESCE(SUM(status='EFEKTIF'),0) ef,
                   COALESCE(SUM(status='TIDAK EFEKTIF'),0) tef, COALESCE(SUM(status IS NULL),0) belum
            FROM disposisi WHERE pegawai_id=?""", (pid,), one=True)
        ctx["upcoming"] = query("""
            SELECT r.*, d.id AS dispo_id, d.status FROM disposisi d JOIN rapat r ON r.id=d.rapat_id
            WHERE d.pegawai_id=? AND r.waktu >= ? ORDER BY r.waktu LIMIT 8""", (pid, now[:10]))
        ctx["perlu"] = query("""
            SELECT r.*, d.id AS dispo_id FROM disposisi d JOIN rapat r ON r.id=d.rapat_id
            WHERE d.pegawai_id=? AND d.status IS NULL AND r.waktu < ? ORDER BY r.waktu DESC""", (pid, now))
    return render_template("dashboard.html", **ctx)


# ---------------------------------------------------------------------------
# Jadwal rapat (admin)
# ---------------------------------------------------------------------------

RAPAT_FIELDS = ["no_dispo_kabag", "no_surat_undangan", "asal_undangan", "perihal",
                "pimpinan_rapat", "tempat", "keterangan"]
RAPAT_LABELS = {
    "no_dispo_kabag": "No. Dispo Kabag", "no_surat_undangan": "No. Surat Undangan",
    "asal_undangan": "Asal Undangan", "perihal": "Perihal (Agenda Rapat)",
    "pimpinan_rapat": "Pimpinan Rapat", "tempat": "Tempat",
}


def rapat_filter_sql():
    q = request.args.get("q", "").strip()
    bulan = request.args.get("bulan", "").strip()
    status = request.args.get("status", "").strip()
    where, args = ["1=1"], []
    if q:
        like = f"%{q}%"
        where.append("""(r.perihal LIKE ? OR r.no_dispo_kabag LIKE ? OR r.no_surat_undangan LIKE ?
                        OR r.asal_undangan LIKE ? OR r.tempat LIKE ? OR r.pimpinan_rapat LIKE ?
                        OR EXISTS (SELECT 1 FROM disposisi d JOIN pegawai p ON p.id=d.pegawai_id
                                   WHERE d.rapat_id=r.id AND p.nama LIKE ?))""")
        args += [like] * 7
    if bulan:
        where.append("substr(r.waktu,1,7) = ?")
        args.append(bulan)
    if status == "belum":
        where.append("EXISTS (SELECT 1 FROM disposisi d WHERE d.rapat_id=r.id AND d.status IS NULL)")
    elif status == "lengkap":
        where.append("NOT EXISTS (SELECT 1 FROM disposisi d WHERE d.rapat_id=r.id AND d.status IS NULL)")
    elif status == "efektif":
        where.append("EXISTS (SELECT 1 FROM disposisi d WHERE d.rapat_id=r.id AND d.status='EFEKTIF')")
    elif status == "tidak":
        where.append("EXISTS (SELECT 1 FROM disposisi d WHERE d.rapat_id=r.id AND d.status='TIDAK EFEKTIF')")
    return " AND ".join(where), args, dict(q=q, bulan=bulan, status=status)


@app.route("/rapat")
@admin_required
def rapat_list():
    where, args, f = rapat_filter_sql()
    rows = attach_dispo(query(f"SELECT r.* FROM rapat r WHERE {where} ORDER BY r.waktu DESC", args))
    return render_template("rapat_list.html", rows=rows, f=f)


def rapat_form_data():
    f = request.form
    data = {k: f.get(k, "").strip() for k in RAPAT_FIELDS}
    errors = [f"{label} wajib diisi." for k, label in RAPAT_LABELS.items() if not data[k]]
    dt = parse_dt(f"{f.get('tanggal', '')}T{f.get('jam', '')}")
    if not dt or not f.get("jam"):
        errors.append("Tanggal dan jam rapat wajib diisi.")
    data["waktu"] = dt.strftime(FMT) if dt else ""
    pegawai_ids = []
    for v in f.getlist("disposisi"):
        try:
            pegawai_ids.append(int(v))
        except ValueError:
            pass
    pegawai_ids = list(dict.fromkeys(pegawai_ids))
    if not pegawai_ids:
        errors.append("Disposisi wajib: pilih minimal satu pegawai yang ditugaskan.")
    return data, pegawai_ids, errors


def pegawai_options(include_ids=()):
    rows = query("SELECT * FROM pegawai WHERE aktif=1 ORDER BY nama")
    if include_ids:
        have = {r["id"] for r in rows}
        missing = [i for i in include_ids if i not in have]
        if missing:
            qs = ",".join("?" * len(missing))
            rows = list(rows) + list(query(f"SELECT * FROM pegawai WHERE id IN ({qs})", missing))
    return rows


@app.route("/rapat/baru", methods=["GET", "POST"])
@admin_required
def rapat_baru():
    if request.method == "POST":
        data, pids, errors = rapat_form_data()
        if errors:
            for e in errors:
                flash(e, "error")
            return render_template("rapat_form.html", r=data, selected=pids, pegawai=pegawai_options(),
                                   tanggal=request.form.get("tanggal"), jam=request.form.get("jam"))
        ts = now_str()
        db = get_db()
        cur = db.execute(f"""INSERT INTO rapat({",".join(RAPAT_FIELDS)}, waktu, created_by, created_at, updated_at)
                             VALUES ({",".join("?" * len(RAPAT_FIELDS))}, ?, ?, ?, ?)""",
                         (*[data[k] for k in RAPAT_FIELDS], data["waktu"], g.user["id"], ts, ts))
        rid = cur.lastrowid
        db.executemany("INSERT INTO disposisi(rapat_id, pegawai_id) VALUES (?, ?)", [(rid, p) for p in pids])
        db.commit()
        flash("Jadwal rapat berhasil ditambahkan.", "success")
        return redirect(url_for("rapat_detail", rid=rid))
    return render_template("rapat_form.html", r={}, selected=[], pegawai=pegawai_options(),
                           tanggal=request.args.get("tanggal", ""), jam="")


@app.route("/rapat/<int:rid>/edit", methods=["GET", "POST"])
@admin_required
def rapat_edit(rid):
    r = query("SELECT * FROM rapat WHERE id=?", (rid,), one=True) or abort(404)
    existing = {d["pegawai_id"]: d for d in query("SELECT * FROM disposisi WHERE rapat_id=?", (rid,))}
    if request.method == "POST":
        data, pids, errors = rapat_form_data()
        if errors:
            for e in errors:
                flash(e, "error")
            data["id"] = rid
            return render_template("rapat_form.html", r=data, selected=pids, pegawai=pegawai_options(pids),
                                   tanggal=request.form.get("tanggal"), jam=request.form.get("jam"),
                                   reported={k for k, v in existing.items() if v["status"]})
        db = get_db()
        db.execute(f"""UPDATE rapat SET {", ".join(f"{k}=?" for k in RAPAT_FIELDS)}, waktu=?, updated_at=?
                       WHERE id=?""", (*[data[k] for k in RAPAT_FIELDS], data["waktu"], now_str(), rid))
        for pid, d in existing.items():
            if pid not in pids:
                remove_file(d["file_path"])
                db.execute("DELETE FROM disposisi WHERE id=?", (d["id"],))
        for pid in pids:
            if pid not in existing:
                db.execute("INSERT INTO disposisi(rapat_id, pegawai_id) VALUES (?, ?)", (rid, pid))
        db.commit()
        flash("Jadwal rapat berhasil diperbarui.", "success")
        return redirect(url_for("rapat_detail", rid=rid))
    dt = parse_dt(r["waktu"])
    return render_template("rapat_form.html", r=dict(r), selected=list(existing.keys()),
                           pegawai=pegawai_options(list(existing.keys())),
                           tanggal=dt.strftime("%Y-%m-%d") if dt else "", jam=dt.strftime("%H:%M") if dt else "",
                           reported={k for k, v in existing.items() if v["status"]})


@app.route("/rapat/<int:rid>/hapus", methods=["POST"])
@admin_required
def rapat_hapus(rid):
    for d in query("SELECT file_path FROM disposisi WHERE rapat_id=?", (rid,)):
        remove_file(d["file_path"])
    execute("DELETE FROM rapat WHERE id=?", (rid,))
    flash("Jadwal rapat dihapus.", "success")
    return redirect(url_for("rapat_list"))


@app.route("/rapat/<int:rid>")
@login_required
def rapat_detail(rid):
    r = query("SELECT * FROM rapat WHERE id=?", (rid,), one=True) or abort(404)
    r = attach_dispo([r])[0]
    if not is_admin():
        mine = [d for d in r["dispo"] if d["pegawai_id"] == g.user["pegawai_id"]]
        if not mine:
            abort(403)
        r["dispo"] = mine  # user hanya melihat disposisi miliknya
    return render_template("rapat_detail.html", r=r)


@app.route("/rapat/export")
@admin_required
def rapat_export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    where, args, f = rapat_filter_sql()
    rows = attach_dispo(query(f"SELECT r.* FROM rapat r WHERE {where} ORDER BY r.waktu", args))
    wb = Workbook()
    ws = wb.active
    ws.title = "Jadwal Rapat"
    judul = f"REKAP JADWAL RAPAT & STATUS KINERJA — {g.settings['nama_instansi']}"
    if f["bulan"]:
        judul += f" — {f_bulan(f['bulan'])}"
    ws.append([judul])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Diekspor: {f_tgljam(now_str())}"])
    ws.append([])
    headers = ["NO", "NO DISPO KABAG", "NO SURAT UNDANGAN", "ASAL UNDANGAN", "PERIHAL (AGENDA RAPAT)",
               "PIMPINAN RAPAT", "TEMPAT", "TANGGAL", "JAM", "DISPOSISI", "STATUS KINERJA", "FILE BUKTI",
               "CATATAN", "KETERANGAN"]
    ws.append(headers)
    thin = Side(style="thin", color="BBBBBB")
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    no = 0
    for r in rows:
        no += 1
        dispos = r["dispo"] or [None]
        for i, d in enumerate(dispos):
            ws.append([
                no if i == 0 else "", r["no_dispo_kabag"], r["no_surat_undangan"], r["asal_undangan"], r["perihal"],
                r["pimpinan_rapat"], r["tempat"], f_tgl(r["waktu"]), f_jam(r["waktu"]),
                d["nama"] if d else "-", (d["status"] or "BELUM LAPOR") if d else "-",
                (d["file_name"] or "-") if d else "-", (d["catatan"] or "") if d else "", r["keterangan"] or "",
            ])
    widths = [5, 22, 22, 22, 40, 22, 24, 24, 8, 24, 16, 26, 34, 26]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        st = row[10].value
        if st == "EFEKTIF":
            row[10].fill = PatternFill("solid", fgColor="D1FAE5")
        elif st == "TIDAK EFEKTIF":
            row[10].fill = PatternFill("solid", fgColor="FEE2E2")
        elif st == "BELUM LAPOR":
            row[10].fill = PatternFill("solid", fgColor="FEF3C7")
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"jadwal-rapat-{f['bulan'] or 'semua'}.xlsx"
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------------------
# Tugas saya & laporan kinerja (user)
# ---------------------------------------------------------------------------

@app.route("/tugas")
@login_required
def tugas_saya():
    pid = g.user["pegawai_id"]
    tab = request.args.get("tab", "semua")
    rows = []
    if pid:
        where = ""
        if tab == "belum":
            where = "AND d.status IS NULL"
        elif tab == "efektif":
            where = "AND d.status='EFEKTIF'"
        elif tab == "tidak":
            where = "AND d.status='TIDAK EFEKTIF'"
        rows = query(f"""SELECT r.*, d.id AS dispo_id, d.status, d.file_name, d.catatan, d.dilaporkan_at
                         FROM disposisi d JOIN rapat r ON r.id=d.rapat_id
                         WHERE d.pegawai_id=? {where} ORDER BY r.waktu DESC""", (pid,))
    return render_template("tugas_list.html", rows=rows, tab=tab)


@app.route("/disposisi/<int:did>/lapor", methods=["GET", "POST"])
@login_required
def lapor(did):
    d = query("""SELECT d.*, p.nama, p.jabatan FROM disposisi d JOIN pegawai p ON p.id=d.pegawai_id
                 WHERE d.id=?""", (did,), one=True) or abort(404)
    if not can_access_dispo(d):
        abort(403)
    r = query("SELECT * FROM rapat WHERE id=?", (d["rapat_id"],), one=True)
    if request.method == "POST":
        status = request.form.get("status", "")
        catatan = request.form.get("catatan", "").strip()
        upload = request.files.get("bukti")
        has_upload = bool(upload and upload.filename)
        if status not in STATUS_LIST:
            flash("Pilih status kinerja: EFEKTIF atau TIDAK EFEKTIF.", "error")
            return render_template("lapor.html", d=d, r=r)
        file_path, file_name = d["file_path"], d["file_name"]
        if status == "EFEKTIF":
            if has_upload:
                ext = upload.filename.rsplit(".", 1)[-1].lower() if "." in upload.filename else ""
                if ext not in ALLOWED_EXT:
                    flash(f"Format file tidak didukung. Gunakan: {', '.join(sorted(ALLOWED_EXT))}.", "error")
                    return render_template("lapor.html", d=d, r=r)
                new_name = f"{uuid.uuid4().hex}.{ext}"
                upload.save(os.path.join(UPLOAD_DIR, new_name))
                remove_file(d["file_path"])
                file_path, file_name = new_name, os.path.basename(upload.filename)[:200]
            elif not d["file_path"]:
                flash("Status EFEKTIF WAJIB melampirkan file bukti (PDF, PNG, JPG, dll).", "error")
                return render_template("lapor.html", d=d, r=r)
        else:  # TIDAK EFEKTIF -> file otomatis kosong
            remove_file(d["file_path"])
            file_path, file_name = None, None
        execute("""UPDATE disposisi SET status=?, file_path=?, file_name=?, catatan=?, dilaporkan_at=?,
                   dilaporkan_oleh=? WHERE id=?""",
                (status, file_path, file_name, catatan, now_str(), g.user["id"], did))
        flash("Laporan status kinerja berhasil disimpan.", "success")
        if is_admin() and d["pegawai_id"] != g.user["pegawai_id"]:
            return redirect(url_for("rapat_detail", rid=d["rapat_id"]))
        return redirect(url_for("tugas_saya"))
    return render_template("lapor.html", d=d, r=r)


@app.route("/berkas/<int:did>")
@login_required
def berkas(did):
    d = query("SELECT * FROM disposisi WHERE id=?", (did,), one=True) or abort(404)
    if not can_access_dispo(d) or not d["file_path"]:
        abort(404 if not d["file_path"] else 403)
    return send_from_directory(UPLOAD_DIR, os.path.basename(d["file_path"]),
                               as_attachment=request.args.get("unduh") == "1", download_name=d["file_name"])


# ---------------------------------------------------------------------------
# Kalender
# ---------------------------------------------------------------------------

@app.route("/kalender")
@login_required
def kalender():
    today = date.today()
    try:
        y = int(request.args.get("y", today.year))
        m = int(request.args.get("m", today.month))
        if not 1 <= m <= 12:
            raise ValueError
    except ValueError:
        y, m = today.year, today.month
    weeks = calendar.Calendar(firstweekday=0).monthdatescalendar(y, m)
    start, end = weeks[0][0], weeks[-1][-1]
    items = {}

    if is_admin():
        rapat = attach_dispo(query("SELECT * FROM rapat WHERE substr(waktu,1,10) BETWEEN ? AND ? ORDER BY waktu",
                                   (start.isoformat(), end.isoformat())))
        for r in rapat:
            cls = "rapat-done" if r["dispo"] and r["jml_belum"] == 0 else "rapat"
            items.setdefault(r["waktu"][:10], []).append(
                dict(cls=cls, jam=f_jam(r["waktu"]), judul=r["perihal"], url=url_for("rapat_detail", rid=r["id"]),
                     info=f"{r['tempat']} • {len(r['dispo'])} pegawai"))
    else:
        rapat = query("""SELECT r.*, d.status, d.id AS dispo_id FROM disposisi d JOIN rapat r ON r.id=d.rapat_id
                         WHERE d.pegawai_id=? AND substr(r.waktu,1,10) BETWEEN ? AND ? ORDER BY r.waktu""",
                      (g.user["pegawai_id"] or -1, start.isoformat(), end.isoformat()))
        for r in rapat:
            cls = {"EFEKTIF": "rapat-done", "TIDAK EFEKTIF": "rapat-no"}.get(r["status"], "rapat")
            items.setdefault(r["waktu"][:10], []).append(
                dict(cls=cls, jam=f_jam(r["waktu"]), judul=r["perihal"], url=url_for("rapat_detail", rid=r["id"]),
                     info=r["tempat"]))

    ev_where, ev_args = visible_events_sql()
    events = query(f"""SELECT e.* FROM events e WHERE {ev_where}
                       AND substr(e.mulai,1,10) <= ? AND substr(coalesce(nullif(e.selesai,''), e.mulai),1,10) >= ?
                       ORDER BY e.mulai""", (*ev_args, end.isoformat(), start.isoformat()))
    for e in events:
        d0 = parse_dt(e["mulai"]).date()
        d1 = (parse_dt(e["selesai"]) or parse_dt(e["mulai"])).date()
        cur = max(d0, start)
        while cur <= min(d1, end):
            items.setdefault(cur.isoformat(), []).append(
                dict(cls=e["jenis"], jam=f_jam(e["mulai"]) if cur == d0 else "", judul=e["judul"],
                     url=url_for("event_list") + f"#ev{e['id']}", info=e["deskripsi"] or ""))
            cur += timedelta(days=1)

    prev_m = (date(y, m, 1) - timedelta(days=1))
    next_m = (date(y, m, 28) + timedelta(days=5))
    return render_template("kalender.html", weeks=weeks, items=items, y=y, m=m, today=today,
                           judul=f"{BULAN[m]} {y}", prev=(prev_m.year, prev_m.month),
                           next=(next_m.year, next_m.month))


# ---------------------------------------------------------------------------
# Event & Reminder
# ---------------------------------------------------------------------------

@app.route("/event")
@login_required
def event_list():
    tab = request.args.get("tab", "mendatang")
    today = date.today().isoformat()
    ev_where, ev_args = visible_events_sql()
    if tab == "lewat":
        cond, order = "substr(coalesce(nullif(e.selesai,''), e.mulai),1,10) < ?", "DESC"
    else:
        cond, order = "substr(coalesce(nullif(e.selesai,''), e.mulai),1,10) >= ?", "ASC"
    rows = query(f"""SELECT e.*, p.nama AS untuk_nama FROM events e LEFT JOIN pegawai p ON p.id=e.untuk_pegawai_id
                     WHERE {ev_where} AND {cond} ORDER BY e.mulai {order}""", (*ev_args, today))
    return render_template("event_list.html", rows=rows, tab=tab)


def event_form_data():
    f = request.form
    errors = []
    judul = f.get("judul", "").strip()
    if not judul:
        errors.append("Judul wajib diisi.")
    tgl = f.get("tanggal", "")
    mulai = parse_dt(f"{tgl}T{f.get('jam') or '00:00'}")
    if not mulai:
        errors.append("Tanggal wajib diisi.")
    selesai = ""
    if f.get("tanggal_selesai") or f.get("jam_selesai"):
        s = parse_dt(f"{f.get('tanggal_selesai') or tgl}T{f.get('jam_selesai') or '23:59'}")
        if s and mulai and s < mulai:
            errors.append("Waktu selesai tidak boleh sebelum waktu mulai.")
        selesai = s.strftime(FMT) if s else ""
    jenis = f.get("jenis") if f.get("jenis") in ("event", "reminder") else "event"
    try:
        ingat = max(0, int(f.get("ingatkan_menit", "0")))
    except ValueError:
        ingat = 0
    untuk = f.get("untuk_pegawai_id") or None
    data = dict(judul=judul, deskripsi=f.get("deskripsi", "").strip(), jenis=jenis,
                mulai=mulai.strftime(FMT) if mulai else "", selesai=selesai, ingatkan_menit=ingat,
                untuk_pegawai_id=int(untuk) if untuk and untuk.isdigit() else None)
    return data, errors


def event_form_ctx(e):
    m, s = parse_dt(e.get("mulai")), parse_dt(e.get("selesai"))
    return dict(e=e, pegawai=pegawai_options(),
                tanggal=m.strftime("%Y-%m-%d") if m else request.args.get("tanggal", ""),
                jam=m.strftime("%H:%M") if m else "",
                tanggal_selesai=s.strftime("%Y-%m-%d") if s else "", jam_selesai=s.strftime("%H:%M") if s else "")


@app.route("/event/baru", methods=["GET", "POST"])
@admin_required
def event_baru():
    if request.method == "POST":
        data, errors = event_form_data()
        if errors:
            for e in errors:
                flash(e, "error")
            return render_template("event_form.html", **event_form_ctx(data))
        execute("""INSERT INTO events(judul,deskripsi,jenis,mulai,selesai,ingatkan_menit,untuk_pegawai_id,
                   created_by,created_at) VALUES (?,?,?,?,?,?,?,?,?)""",
                (*data.values(), g.user["id"], now_str()))
        flash("Event/reminder berhasil ditambahkan.", "success")
        return redirect(url_for("event_list"))
    return render_template("event_form.html", **event_form_ctx({"jenis": request.args.get("jenis", "event"),
                                                                 "ingatkan_menit": 60}))


@app.route("/event/<int:eid>/edit", methods=["GET", "POST"])
@admin_required
def event_edit(eid):
    e = query("SELECT * FROM events WHERE id=?", (eid,), one=True) or abort(404)
    if request.method == "POST":
        data, errors = event_form_data()
        if errors:
            for er in errors:
                flash(er, "error")
            data["id"] = eid
            return render_template("event_form.html", **event_form_ctx(data))
        execute("""UPDATE events SET judul=?,deskripsi=?,jenis=?,mulai=?,selesai=?,ingatkan_menit=?,
                   untuk_pegawai_id=? WHERE id=?""", (*data.values(), eid))
        flash("Event/reminder diperbarui.", "success")
        return redirect(url_for("event_list"))
    return render_template("event_form.html", **event_form_ctx(dict(e)))


@app.route("/event/<int:eid>/hapus", methods=["POST"])
@admin_required
def event_hapus(eid):
    execute("DELETE FROM events WHERE id=?", (eid,))
    flash("Event/reminder dihapus.", "success")
    return redirect(url_for("event_list"))


@app.route("/api/reminders")
@login_required
def api_reminders():
    """Dipanggil berkala oleh browser untuk memunculkan notifikasi pengingat."""
    now = datetime.now()
    out = []
    ev_where, ev_args = visible_events_sql()
    horizon = (now + timedelta(days=8)).strftime(FMT)
    for e in query(f"""SELECT e.* FROM events e WHERE {ev_where} AND e.ingatkan_menit > 0
                       AND e.mulai >= ? AND e.mulai <= ?""",
                   (*ev_args, (now - timedelta(minutes=1)).strftime(FMT), horizon)):
        start = parse_dt(e["mulai"])
        if start - timedelta(minutes=e["ingatkan_menit"]) <= now <= start:
            out.append(dict(key=f"e{e['id']}-{e['mulai']}", jenis=e["jenis"].capitalize(), judul=e["judul"],
                            waktu=f_tgljam(e["mulai"]), sisa=f_relatif(e["mulai"]), url=url_for("event_list")))
    try:
        menit = int(g.settings.get("pengingat_rapat_menit", "60"))
    except ValueError:
        menit = 60
    if g.user["pegawai_id"] and menit > 0:
        for r in query("""SELECT r.* FROM disposisi d JOIN rapat r ON r.id=d.rapat_id
                          WHERE d.pegawai_id=? AND r.waktu >= ? AND r.waktu <= ?""",
                       (g.user["pegawai_id"], now.strftime(FMT), (now + timedelta(minutes=menit)).strftime(FMT))):
            out.append(dict(key=f"r{r['id']}-{r['waktu']}", jenis="Rapat", judul=r["perihal"],
                            waktu=f"{f_tgljam(r['waktu'])} • {r['tempat']}", sisa=f_relatif(r["waktu"]),
                            url=url_for("rapat_detail", rid=r["id"])))
    return jsonify(out)


# ---------------------------------------------------------------------------
# Data pegawai (admin)
# ---------------------------------------------------------------------------
PEGAWAI_FIELDS = ["nip", "nama", "jabatan", "unit", "no_hp", "email"]


@app.route("/pegawai")
@admin_required
def pegawai_list():
    q = request.args.get("q", "").strip()
    where, args = "1=1", []
    if q:
        where = "(p.nama LIKE ? OR p.nip LIKE ? OR p.jabatan LIKE ? OR p.unit LIKE ?)"
        args = [f"%{q}%"] * 4
    rows = query(f"""SELECT p.*, u.username, u.id AS user_id,
                     (SELECT COUNT(*) FROM disposisi d WHERE d.pegawai_id=p.id) AS jml_tugas
                     FROM pegawai p LEFT JOIN users u ON u.pegawai_id=p.id
                     WHERE {where} ORDER BY p.aktif DESC, p.nama""", args)
    return render_template("pegawai_list.html", rows=rows, q=q)


def pegawai_save(pid=None):
    data = {k: request.form.get(k, "").strip() for k in PEGAWAI_FIELDS}
    data["aktif"] = 1 if request.form.get("aktif") else 0
    if not data["nama"]:
        flash("Nama pegawai wajib diisi.", "error")
        return None, data
    if pid:
        execute(f"UPDATE pegawai SET {', '.join(f'{k}=?' for k in data)} WHERE id=?", (*data.values(), pid))
    else:
        pid = execute(f"INSERT INTO pegawai({', '.join(data)}, created_at) VALUES ({', '.join('?' * len(data))}, ?)",
                      (*data.values(), now_str()))
    return pid, data


@app.route("/pegawai/baru", methods=["GET", "POST"])
@admin_required
def pegawai_baru():
    if request.method == "POST":
        pid, data = pegawai_save()
        if pid:
            flash("Data pegawai ditambahkan.", "success")
            if request.form.get("buat_akun"):
                return redirect(url_for("pengguna_baru", pegawai_id=pid))
            return redirect(url_for("pegawai_list"))
        return render_template("pegawai_form.html", p=data)
    return render_template("pegawai_form.html", p={"aktif": 1})


@app.route("/pegawai/<int:pid>/edit", methods=["GET", "POST"])
@admin_required
def pegawai_edit(pid):
    p = query("SELECT * FROM pegawai WHERE id=?", (pid,), one=True) or abort(404)
    if request.method == "POST":
        ok, data = pegawai_save(pid)
        if ok:
            flash("Data pegawai diperbarui.", "success")
            return redirect(url_for("pegawai_list"))
        data["id"] = pid
        return render_template("pegawai_form.html", p=data)
    return render_template("pegawai_form.html", p=dict(p))


@app.route("/pegawai/<int:pid>/hapus", methods=["POST"])
@admin_required
def pegawai_hapus(pid):
    n = query("SELECT COUNT(*) c FROM disposisi WHERE pegawai_id=?", (pid,), one=True)["c"]
    if n:
        flash(f"Pegawai memiliki {n} riwayat tugas sehingga tidak dapat dihapus. "
              "Nonaktifkan saja agar riwayat kinerja tetap tersimpan.", "error")
    else:
        execute("DELETE FROM pegawai WHERE id=?", (pid,))
        flash("Data pegawai dihapus.", "success")
    return redirect(url_for("pegawai_list"))


# ---------------------------------------------------------------------------
# Akses pengguna (admin)
# ---------------------------------------------------------------------------

@app.route("/pengguna")
@admin_required
def pengguna_list():
    rows = query("""SELECT u.*, p.nama AS pegawai_nama, p.jabatan FROM users u
                    LEFT JOIN pegawai p ON p.id=u.pegawai_id ORDER BY u.role, u.username""")
    return render_template("pengguna_list.html", rows=rows)


def pegawai_tanpa_akun(include=None):
    return query("""SELECT * FROM pegawai p WHERE p.aktif=1 AND (p.id=? OR NOT EXISTS
                    (SELECT 1 FROM users u WHERE u.pegawai_id=p.id)) ORDER BY p.nama""", (include or -1,))


def pengguna_save(uid=None):
    f = request.form
    username = f.get("username", "").strip().lower()
    password = f.get("password", "")
    role = f.get("role") if f.get("role") in ("admin", "user") else "user"
    pegawai_id = int(f["pegawai_id"]) if f.get("pegawai_id", "").isdigit() else None
    aktif = 1 if f.get("aktif") else 0
    errors = []
    if not username or not username.replace(".", "").replace("_", "").isalnum():
        errors.append("Username wajib diisi (huruf/angka, boleh titik atau garis bawah).")
    elif query("SELECT 1 FROM users WHERE lower(username)=? AND id != ?", (username, uid or -1), one=True):
        errors.append("Username sudah dipakai.")
    if not uid and len(password) < 6:
        errors.append("Password minimal 6 karakter.")
    if uid and password and len(password) < 6:
        errors.append("Password baru minimal 6 karakter.")
    if role == "user" and not pegawai_id:
        errors.append("Akun User wajib dihubungkan ke data pegawai agar bisa menerima tugas.")
    if uid and uid == g.user["id"] and (role != "admin" or not aktif):
        errors.append("Anda tidak dapat menurunkan peran / menonaktifkan akun Anda sendiri.")
    data = dict(username=username, role=role, pegawai_id=pegawai_id, aktif=aktif)
    if errors:
        for e in errors:
            flash(e, "error")
        return None, data
    if uid:
        execute("UPDATE users SET username=?, role=?, pegawai_id=?, aktif=? WHERE id=?",
                (username, role, pegawai_id, aktif, uid))
        if password:
            execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(password), uid))
    else:
        uid = execute("""INSERT INTO users(username,password_hash,role,pegawai_id,aktif,created_at)
                         VALUES (?,?,?,?,?,?)""",
                      (username, generate_password_hash(password), role, pegawai_id, aktif, now_str()))
    return uid, data


@app.route("/pengguna/baru", methods=["GET", "POST"])
@admin_required
def pengguna_baru():
    if request.method == "POST":
        uid, data = pengguna_save()
        if uid:
            flash("Akun pengguna dibuat.", "success")
            return redirect(url_for("pengguna_list"))
        return render_template("pengguna_form.html", u=data, pegawai=pegawai_tanpa_akun(data["pegawai_id"]))
    pid = request.args.get("pegawai_id", type=int)
    return render_template("pengguna_form.html", u={"role": "user", "aktif": 1, "pegawai_id": pid},
                           pegawai=pegawai_tanpa_akun(pid))


@app.route("/pengguna/<int:uid>/edit", methods=["GET", "POST"])
@admin_required
def pengguna_edit(uid):
    u = query("SELECT * FROM users WHERE id=?", (uid,), one=True) or abort(404)
    if request.method == "POST":
        ok, data = pengguna_save(uid)
        if ok:
            flash("Akun pengguna diperbarui.", "success")
            return redirect(url_for("pengguna_list"))
        data["id"] = uid
        return render_template("pengguna_form.html", u=data, pegawai=pegawai_tanpa_akun(data["pegawai_id"]))
    return render_template("pengguna_form.html", u=dict(u), pegawai=pegawai_tanpa_akun(u["pegawai_id"]))


@app.route("/pengguna/<int:uid>/hapus", methods=["POST"])
@admin_required
def pengguna_hapus(uid):
    if uid == g.user["id"]:
        flash("Anda tidak dapat menghapus akun sendiri.", "error")
    else:
        execute("DELETE FROM users WHERE id=?", (uid,))
        flash("Akun pengguna dihapus.", "success")
    return redirect(url_for("pengguna_list"))


# ---------------------------------------------------------------------------
# Rekap kinerja (admin)
# ---------------------------------------------------------------------------

def rekap_data(bulan):
    cond, args = ("WHERE substr(r.waktu,1,7)=?", [bulan]) if bulan else ("", [])
    rows = query(f"""
        SELECT p.id, p.nip, p.nama, p.jabatan, p.unit,
               COUNT(x.id) total,
               COALESCE(SUM(x.status='EFEKTIF'),0) ef,
               COALESCE(SUM(x.status='TIDAK EFEKTIF'),0) tef,
               COALESCE(SUM(x.id IS NOT NULL AND x.status IS NULL),0) belum
        FROM pegawai p
        LEFT JOIN (SELECT d.* FROM disposisi d JOIN rapat r ON r.id=d.rapat_id {cond}) x ON x.pegawai_id=p.id
        WHERE p.aktif=1 OR x.id IS NOT NULL
        GROUP BY p.id ORDER BY total DESC, p.nama""", args)
    out = []
    for r in rows:
        r = dict(r)
        r["persen"] = round(r["ef"] * 100 / r["total"]) if r["total"] else 0
        out.append(r)
    return out


@app.route("/rekap")
@admin_required
def rekap():
    bulan = request.args.get("bulan", date.today().strftime("%Y-%m"))
    rows = rekap_data(bulan)
    tot = dict(total=sum(r["total"] for r in rows), ef=sum(r["ef"] for r in rows),
               tef=sum(r["tef"] for r in rows), belum=sum(r["belum"] for r in rows))
    tot["persen"] = round(tot["ef"] * 100 / tot["total"]) if tot["total"] else 0
    return render_template("rekap.html", rows=rows, bulan=bulan, tot=tot)


@app.route("/rekap/export")
@admin_required
def rekap_export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    bulan = request.args.get("bulan", "")
    rows = rekap_data(bulan)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Kinerja"
    ws.append([f"REKAP KINERJA PEGAWAI — {g.settings['nama_instansi']} — {f_bulan(bulan) if bulan else 'Semua Periode'}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["NO", "NIP", "NAMA", "JABATAN", "UNIT", "TOTAL TUGAS", "EFEKTIF", "TIDAK EFEKTIF",
               "BELUM LAPOR", "% EFEKTIF"])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
        c.alignment = Alignment(horizontal="center")
    for i, r in enumerate(rows, 1):
        ws.append([i, r["nip"], r["nama"], r["jabatan"], r["unit"], r["total"], r["ef"], r["tef"], r["belum"],
                   f"{r['persen']}%"])
    for col, w in zip("ABCDEFGHIJ", [5, 22, 26, 24, 20, 12, 10, 14, 12, 10]):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"rekap-kinerja-{bulan or 'semua'}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------------------
# Pengaturan sistem (admin)
# ---------------------------------------------------------------------------

@app.route("/pengaturan", methods=["GET", "POST"])
@admin_required
def pengaturan():
    if request.method == "POST":
        for k in DEFAULT_SETTINGS:
            v = request.form.get(k, "").strip()
            if k == "pengingat_rapat_menit":
                v = str(max(0, int(v))) if v.isdigit() else "60"
            if v:
                execute("INSERT INTO settings(key,value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                        (k, v))
        flash("Pengaturan disimpan.", "success")
        return redirect(url_for("pengaturan"))
    info = dict(
        pegawai=query("SELECT COUNT(*) c FROM pegawai", one=True)["c"],
        users=query("SELECT COUNT(*) c FROM users", one=True)["c"],
        rapat=query("SELECT COUNT(*) c FROM rapat", one=True)["c"],
        events=query("SELECT COUNT(*) c FROM events", one=True)["c"],
        files=len(os.listdir(UPLOAD_DIR)),
    )
    return render_template("pengaturan.html", info=info)


@app.route("/pengaturan/backup")
@admin_required
def backup_db():
    return send_file(DB_PATH, as_attachment=True,
                     download_name=f"backup-monitoring-{datetime.now().strftime('%Y%m%d-%H%M')}.db")


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)),
            debug=os.environ.get("FLASK_DEBUG") == "1")
